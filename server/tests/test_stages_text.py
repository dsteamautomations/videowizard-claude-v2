"""Text stages T1–T7: intake → brief → script through the vw plugin against the live test server."""
import json
from io import BytesIO

import pytest

from conftest import rid, seed_session
from conftest import connected, vw  # noqa: F401

SOURCE = """Scene 1: The yard

Narration: Earth blocks are made from the soil beneath our feet.

Duration: 30 seconds

Visuals
- Close-up of finished blocks

Narration: Each block is pressed, not fired, so it keeps the carbon in the ground.
"""


def start_run(home, database):
    connected(home, database)
    code, out, err = vw(home, 'start', '--client', 'VK', '--need-gb', '0')
    assert code == 0, err
    return out['code'], home['runs'] / out['code']


def write_proposal(ws, stage, data):
    (ws / 'out' / f'{stage}.json').write_text(json.dumps(data), encoding='utf-8')


def test_t7_intake_brief_script_end_to_end(home, database):
    code, ws = start_run(home, database)
    (ws / 'in' / 'source.txt').write_text(SOURCE, encoding='utf-8')
    rc, out, err = vw(home, 'stage', 'intake', 'prepare')
    assert rc == 0, err
    assert '{numbered_script}' not in out['prompt'] and 'Scene 1: The yard' not in out['prompt']   # design/34 1.1: the numbered script is a file
    assert (ws / 'work' / 'intake-numbered.md').read_text(encoding='utf-8').startswith('1: Scene 1: The yard\n') and out['lines'] == 11
    assert out['prompt'].startswith('VW-PROCEDURE') and out['not_shareable'] and out['earlier_titles'] == [] and out['prepared_at']
    assert out['needs_readable_copy'] is False and out['output_file'] == 'out/intake.json'
    write_proposal(ws, 'intake', {'remove': [{'line': 1, 'kind': 'scene_heading', 'why': 'heading'}, {'line': 5, 'kind': 'timing', 'why': 'duration'},
                                             {'line': 7, 'kind': 'section_label', 'why': 'label'}, {'line': 8, 'kind': 'visual_note', 'why': 'seen not said'}],
                                  'titles': ['Earth Blocks - Documentary Explainer', 'Pressed Earth - Explainer'], 'notes': 'removed plan lines'})
    # the operator picks the second suggestion; the client is the run's directory client
    rc, out, _ = vw(home, 'stage', 'intake', 'record', '--quote', 'yes, approved, the second one')
    assert rc == 2 and '--title' in out['message']   # design/34 1.2: the operator's choice is required
    rc, out, err = vw(home, 'stage', 'intake', 'record', '--title', 'Pressed Earth - Explainer', '--quote', 'yes, approved, the second one')
    assert rc == 0, (out, err)
    assert out['summary']['title'] == 'Pressed Earth - Explainer' and out['summary']['client'] == 'VK' and out['summary']['title_source'] == 'suggested'
    assert out['summary']['prepared_at'] and out['summary']['recorded_at'] and out['next']['stage'] == 'brief' and 'Next: brief' in err   # 1.15, 1.19
    import httpx
    tok = json.loads((home['home'] / 'session.json').read_text())['token']
    run = httpx.get(f"{home['env']['VW2_SERVER_URL']}/api/runs/{out['run_id']}", headers={'Authorization': 'Bearer ' + tok}).json()
    assert run['title'] == 'Pressed Earth - Explainer' and run['client'] == 'VK' and run['client_id'] == 'c-vk'
    cleaned = (ws / 'out' / 'source.txt').read_text(encoding='utf-8')
    assert 'Scene 1' not in cleaned and 'Duration' not in cleaned and 'Close-up' not in cleaned
    assert 'Narration: Earth blocks are made' in cleaned and 'keeps the carbon' in cleaned
    assert vw(home, 'next')[1]['stage'] == 'brief'

    rc, out, err = vw(home, 'stage', 'brief', 'prepare')
    assert rc == 0 and 'keeps the carbon' in out['inputs']['script'] and out['inputs']['title'] == 'Pressed Earth - Explainer'
    write_proposal(ws, 'brief', {'purpose': 'Understand why earth blocks matter', 'audience': 'self-builders (recommendation)', 'message': 'Pressed, not fired',
                                 'tone': 'calm', 'ending': 'invite to visit the yard', 'constraints': '', 'reasons': {'audience': 'subject suggests it'}, 'questions': [],
                                 'production': {'field_speech_languages': ['ta'], 'narration_language': 'en', 'brand_and_safety_rules': ['PPE required on site']}})
    # a required question without the operator's answer blocks record; the answer is stored verbatim
    proposal = json.loads((ws / 'out' / 'brief.json').read_text(encoding='utf-8'))
    proposal['questions'] = [{'id': 'q1', 'question': 'Sales ending or awareness only?', 'why': 'changes the ending', 'recommendation': 'awareness', 'required': True}]
    write_proposal(ws, 'brief', proposal)
    rc, out, _ = vw(home, 'stage', 'brief', 'record', '--quote', 'yes')
    assert rc == 2 and 'q1' in out['message']
    proposal['answers'] = {'q1': 'awareness only, no sales'}
    write_proposal(ws, 'brief', proposal)
    rc, out, err = vw(home, 'stage', 'brief', 'record', '--quote', 'yes')
    assert rc == 0, (out, err)
    assert out['summary']['answers'] == {'q1': 'awareness only, no sales'} and out['summary']['questions'] == 1
    prod = out['summary']['production']
    assert prod['process_stages'] == [] and 'process_stages' in prod['defaults_applied']   # chapters: empty unless the script gives them (design/37 2.1)
    assert prod['show_speaker_video'] is False and prod['target_runtime_s'] is None and prod['field_speech_languages'] == ['ta'] and 'field_speech_languages' not in prod['defaults_applied']
    assert vw(home, 'next')[1]['stage'] == 'script'

    rc, out, err = vw(home, 'stage', 'script', 'prepare')
    assert rc == 0 and out['inputs']['brief']['message'] == 'Pressed, not fired' and 'keeps the carbon' in out['inputs']['script']
    write_proposal(ws, 'script', {'spoken_text': 'Earth blocks are made from the soil beneath our feet.\n\nEach block is pressed, not fired.',
                                  'notes': '', 'assessment': 'fits', 'changes': [{'before': 'keeps the carbon', 'after': 'keeps the carbon in the ground', 'reason': 'brief tone'}]})
    rc, out, err = vw(home, 'stage', 'script', 'record', '--quote', 'approved')
    assert rc == 0, err
    assert out['summary']['passages'] == 2 and out['summary']['beats'] == 2 and len(out['summary']['content_digest']) == 64
    script = json.loads((ws / 'out' / 'script.json').read_text(encoding='utf-8'))
    assert [b['beat_id'] for b in script['beats']] == ['BEAT_01', 'BEAT_02'] and script['beats'][1]['passage_id'] == 'P002' and script['beats'][1]['text'].startswith('Each block')
    # beat ids are stable: recording the same script again yields the same beats
    assert vw(home, 'revise', '--stage', 'script')[0] == 0
    assert vw(home, 'stage', 'script', 'record', '--quote', 'approved again')[0] == 0
    assert json.loads((ws / 'out' / 'script.json').read_text(encoding='utf-8'))['beats'] == script['beats']
    assert list(script['passages']) == ['P001', 'P002']
    nxt = vw(home, 'next')[1]
    assert nxt['stage'] == 'voice' and nxt['paid'] == 'elevenlabs'


def test_t1_docx_tables_included_images_not(home, database):
    from docx import Document
    code, ws = start_run(home, database)
    doc = Document()
    doc.add_paragraph('Narration: hello from a paragraph')
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = 'cell one', 'cell two'
    doc.save(ws / 'in' / 'script.docx')
    rc, out, err = vw(home, 'stage', 'intake', 'prepare')
    assert rc == 0, err
    assert 'hello from a paragraph' in out['prompt'] and 'cell one\tcell two' in out['prompt']
    assert any('images' in w for w in out['warnings'])


def test_t2_pdf_unreadable_page(home, database):
    from pypdf import PdfWriter
    code, ws = start_run(home, database)
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    with open(ws / 'in' / 'script.pdf', 'wb') as f:
        w.write(f)
    rc, out, err = vw(home, 'stage', 'intake', 'prepare')
    assert rc == 0, err
    assert out['needs_readable_copy'] is True and any('pages 1' in w for w in out['warnings'])


def test_t3_source_too_large(home, database):
    code, ws = start_run(home, database)
    (ws / 'in' / 'big.txt').write_bytes(b'x' * (20 * 1024 * 1024 + 1))
    rc, out, err = vw(home, 'stage', 'intake', 'prepare')
    assert rc == 2 and out['action'] == 'refused' and '20 MB' in out['message']


def test_t4_malformed_proposals_refused(home, database):
    code, ws = start_run(home, database)
    (ws / 'in' / 'source.txt').write_text(SOURCE, encoding='utf-8')
    vw(home, 'stage', 'intake', 'prepare')
    write_proposal(ws, 'intake', {'remove': [{'line': 2, 'kind': 'timing', 'why': 'blank'}], 'titles': ['T'], 'notes': ''})
    rc, out, _ = vw(home, 'stage', 'intake', 'record', '--quote', 'yes')
    assert rc == 2 and out['action'] == 'refused' and 'Blank' in out['message']
    write_proposal(ws, 'intake', {'remove': [{'line': 99, 'kind': 'timing', 'why': 'x'}], 'titles': ['T'], 'notes': ''})
    assert vw(home, 'stage', 'intake', 'record', '--quote', 'yes')[0] == 2
    write_proposal(ws, 'intake', {'remove': [], 'titles': [], 'notes': ''})
    assert vw(home, 'stage', 'intake', 'record', '--quote', 'yes')[0] == 2  # a title is required
    write_proposal(ws, 'intake', {'remove': [], 'titles': ['Earth'], 'notes': '', 'client': 'VK'})
    assert vw(home, 'stage', 'intake', 'record', '--quote', 'yes')[0] == 2   # no --title
    write_proposal(ws, 'intake', {'remove': [], 'titles': [42], 'notes': ''})
    rc, out, _ = vw(home, 'stage', 'intake', 'record', '--title', 'Earth', '--quote', 'yes')
    assert rc == 2 and out['action'] == 'refused' and 'titles' in out['message']   # a non-string is refused, never a traceback (design/34 1.4)
    write_proposal(ws, 'intake', {'remove': [], 'titles': ['Earth'], 'notes': '', 'client': 'VK'})
    assert vw(home, 'stage', 'intake', 'record', '--title', 'Earth', '--quote', 'yes')[0] == 0
    write_proposal(ws, 'brief', {'purpose': 'p', 'audience': '', 'message': 'm', 'tone': 't', 'ending': 'e'})
    rc, out, _ = vw(home, 'stage', 'brief', 'record', '--quote', 'yes')
    assert rc == 2 and 'audience' in out['message']
    write_proposal(ws, 'brief', {'purpose': 'p', 'audience': 'a', 'message': 'm', 'tone': 't', 'ending': 'e'})
    rc, out, _ = vw(home, 'stage', 'brief', 'record', '--quote', 'yes')
    assert rc == 2 and 'every film' in out['message']   # languages are asked on every film (design/37 2.8)
    write_proposal(ws, 'brief', {'purpose': 'p', 'audience': 'a', 'message': 'm', 'tone': 't', 'ending': 'e', 'production': {'field_speech_languages': ['ta'], 'narration_language': 'en',
                                 'process_stages': ['Stage_1_Setup', 'Stage_2_Preparation', 'Stage_3_Execution', 'Stage_4_Result']}})
    rc, out, _ = vw(home, 'stage', 'brief', 'record', '--quote', 'yes')
    assert rc == 2 and 'chapters' in out['message']   # the generic names are refused (design/37 2.1)
    write_proposal(ws, 'brief', {'purpose': 'p', 'audience': 'a', 'message': 'm', 'tone': 't', 'ending': 'e', 'production': {'field_speech_languages': ['ta'], 'narration_language': 'en'}})
    assert vw(home, 'stage', 'brief', 'record', '--quote', 'yes')[0] == 0
    write_proposal(ws, 'script', {'spoken_text': '   ', 'notes': '', 'assessment': '', 'changes': []})
    rc, out, _ = vw(home, 'stage', 'script', 'record', '--quote', 'yes')
    assert rc == 2 and 'spoken_text' in out['message']


def test_t5_revise_brief_after_script_is_free(home, database):
    code, ws = start_run(home, database)
    (ws / 'in' / 'source.txt').write_text('Just narration.\n', encoding='utf-8')
    vw(home, 'stage', 'intake', 'prepare')
    write_proposal(ws, 'intake', {'remove': [], 'titles': ['Only'], 'notes': '', 'client': 'VK'})
    assert vw(home, 'stage', 'intake', 'record', '--title', 'Only', '--quote', 'yes')[0] == 0
    write_proposal(ws, 'brief', {'purpose': 'p', 'audience': 'a', 'message': 'm', 'tone': 't', 'ending': 'e', 'production': {'field_speech_languages': ['ta'], 'narration_language': 'en'}})
    assert vw(home, 'stage', 'brief', 'record', '--quote', 'yes')[0] == 0
    write_proposal(ws, 'script', {'spoken_text': 'Just narration.', 'notes': '', 'assessment': '', 'changes': []})
    assert vw(home, 'stage', 'script', 'record', '--quote', 'yes')[0] == 0
    rc, out, _ = vw(home, 'revise', '--stage', 'brief')
    assert rc == 0 and out['reopened'] == ['brief', 'script', 'voice', 'media', 'analysis', 'options', 'master', 'first_cut', 'visuals', 'assembly', 'delivery']  # version 2 stage order
    assert vw(home, 'next')[1]['stage'] == 'brief'
    assert vw(home, 'view', '--stage', 'script')[1]['outputs'][0]['stale'] is True


def test_t6_procedures_executor_only(client, database):
    assert client.get('/api/procedures/1/intake', headers=seed_session(database, 'x', kind='browser')).status_code == 403
    r = client.get('/api/procedures/1/intake', headers=seed_session(database, 'x', kind='executor'))
    assert r.status_code == 200 and 'YOU NEVER REWRITE A LINE' in r.json()['prompt'] and r.json()['output_file'] == 'intake.json'
    assert client.get('/api/procedures/1/nope', headers=seed_session(database, 'y', kind='executor')).status_code == 422
    assert client.get('/api/procedures/9/intake', headers=seed_session(database, 'z', kind='executor')).status_code == 404


def test_t8_two_documents_refuse_until_named(home, database):
    """design/34 1.3"""
    code, ws = start_run(home, database)
    (ws / 'in' / 'source.txt').write_text(SOURCE, encoding='utf-8')
    (ws / 'in' / 'brief.md').write_text('# not the script\n', encoding='utf-8')
    rc, out, _ = vw(home, 'stage', 'intake', 'prepare')
    assert rc == 2 and out['action'] == 'refused' and 'brief.md' in out['message'] and 'source.txt' in out['message']
    rc, out, _ = vw(home, 'stage', 'intake', 'prepare', '--file', 'source.txt')
    assert rc == 0 and out['source_file'] == 'source.txt'


def test_t9_source_replaced_after_prepare(home, database):
    """design/34 1.8"""
    code, ws = start_run(home, database)
    (ws / 'in' / 'source.txt').write_text(SOURCE, encoding='utf-8')
    assert vw(home, 'stage', 'intake', 'prepare')[0] == 0
    write_proposal(ws, 'intake', {'remove': [], 'titles': ['Earth'], 'notes': ''})
    (ws / 'in' / 'source.txt').write_text(SOURCE + 'Narration: one more line.\n', encoding='utf-8')
    rc, out, _ = vw(home, 'stage', 'intake', 'record', '--title', 'Earth', '--quote', 'yes')
    assert rc == 2 and 'changed' in out['message'] and 'prepare' in out['message']
    assert vw(home, 'stage', 'intake', 'prepare')[0] == 0
    assert vw(home, 'stage', 'intake', 'record', '--title', 'Earth', '--quote', 'yes')[0] == 0


def test_t10_cleaned_script_edited_before_brief(home, database):
    """design/34 1.14"""
    code, ws = start_run(home, database)
    (ws / 'in' / 'source.txt').write_text(SOURCE, encoding='utf-8')
    vw(home, 'stage', 'intake', 'prepare')
    write_proposal(ws, 'intake', {'remove': [], 'titles': ['Earth'], 'notes': ''})
    assert vw(home, 'stage', 'intake', 'record', '--title', 'Earth', '--quote', 'yes')[0] == 0
    cleaned = ws / 'out' / 'source.txt'
    cleaned.write_text(cleaned.read_text(encoding='utf-8').replace('pressed', 'PRESSED'), encoding='utf-8')
    rc, out, _ = vw(home, 'stage', 'brief', 'prepare')
    assert rc == 2 and 'differs from the recorded one' in out['message']


def test_t11_earlier_titles_reach_the_next_film(home, database):
    """design/34 1.5 and 1.17: the second film for a client sees the first one's title; a partial name resolves."""
    code, ws = start_run(home, database)
    (ws / 'in' / 'source.txt').write_text(SOURCE, encoding='utf-8')
    vw(home, 'stage', 'intake', 'prepare')
    write_proposal(ws, 'intake', {'remove': [], 'titles': ['Earth Blocks - Documentary Explainer'], 'notes': ''})
    assert vw(home, 'stage', 'intake', 'record', '--title', 'Earth Blocks - Documentary Explainer', '--quote', 'yes')[0] == 0
    rc, out, _ = vw(home, 'start', '--client', 'vk', '--need-gb', '0')
    assert rc == 0 and out['client'] == 'VK' and out['client_runs'] == [{'code': code, 'title': 'Earth Blocks - Documentary Explainer'}]
    ws2 = home['runs'] / out['code']
    (ws2 / 'in' / 'source.txt').write_text(SOURCE, encoding='utf-8')
    rc, out, _ = vw(home, 'stage', 'intake', 'prepare')
    assert rc == 0 and f'{code}: Earth Blocks - Documentary Explainer' in out['prompt'] and out['earlier_titles'][0]['code'] == code


def test_t12_start_with_script_copies_it(home, database, tmp_path):
    """design/34 1.6"""
    connected(home, database)
    src = tmp_path / 'client-script.txt'
    src.write_text(SOURCE, encoding='utf-8')
    rc, out, _ = vw(home, 'start', '--client', 'VK', '--need-gb', '0', '--script', str(src))
    assert rc == 0 and out['script_file'] == 'client-script.txt' and out['next']['stage'] == 'intake'
    assert (home['runs'] / out['code'] / 'in' / 'client-script.txt').read_text(encoding='utf-8') == SOURCE
    rc, out, _ = vw(home, 'start', '--client', 'VK', '--need-gb', '0', '--script', str(tmp_path / 'missing.docx'))
    assert rc == 2 and 'not created' in out['message']


def test_t13_last_brief_and_client_rules_reach_the_next_film(home, database):
    """design/37 2.2, 2.3: a second film for the client starts with the first film's brief answers and the learned rules."""
    import psycopg
    code, ws = start_run(home, database)
    (ws / 'in' / 'source.txt').write_text(SOURCE, encoding='utf-8')
    vw(home, 'stage', 'intake', 'prepare')
    write_proposal(ws, 'intake', {'remove': [], 'titles': ['Earth'], 'notes': ''})
    assert vw(home, 'stage', 'intake', 'record', '--title', 'Earth', '--quote', 'yes')[0] == 0
    rc, out, _ = vw(home, 'stage', 'brief', 'prepare')
    assert rc == 0 and out['inputs']['last_brief'] is None and out['inputs']['client_rules'] == [] and out['prompt'].count('\nnone\n') == 2
    write_proposal(ws, 'brief', {'purpose': 'p', 'audience': 'a', 'message': 'm', 'tone': 't', 'ending': 'e', 'constraints': '',
                                 'questions': [{'id': 'brand-safety', 'question': 'Any brand or safety rules?', 'why': 'w', 'recommendation': '', 'required': True}],
                                 'answers': {'brand-safety': 'PPE must be worn'},
                                 'production': {'field_speech_languages': ['ta'], 'narration_language': 'en', 'brand_and_safety_rules': ['PPE required on site'], 'process_stages': ['Plinth', 'Curing']}})
    assert vw(home, 'stage', 'brief', 'record', '--quote', 'yes')[0] == 0
    with psycopg.connect(database) as conn:
        conn.execute("insert into vw2_client_rules(client_id, rule, source_run) values ('c-vk', 'Rejected option #2: no unfinished walls', %s)", (code,))
    rc, out, _ = vw(home, 'start', '--client', 'VK', '--need-gb', '0')
    assert rc == 0, out
    lb = out['last_brief']
    assert lb['run'] == code and lb['production']['process_stages'] == ['Plinth', 'Curing'] and lb['answers'] == {'brand-safety': 'PPE must be worn'}
    assert lb['questions'] == [{'id': 'brand-safety', 'question': 'Any brand or safety rules?'}]
    assert [r['rule'] for r in out['client_rules']] == ['Rejected option #2: no unfinished walls']
    ws2 = home['runs'] / out['code']
    (ws2 / 'in' / 'source.txt').write_text(SOURCE, encoding='utf-8')
    vw(home, 'stage', 'intake', 'prepare')
    write_proposal(ws2, 'intake', {'remove': [], 'titles': ['Earth v2'], 'notes': ''})
    assert vw(home, 'stage', 'intake', 'record', '--title', 'Earth v2', '--quote', 'yes')[0] == 0
    rc, out, _ = vw(home, 'stage', 'brief', 'prepare')
    assert rc == 0 and f'from run {code}:' in out['prompt'] and 'process_stages: Plinth, Curing' in out['prompt']
    assert 'answer to brand-safety (Any brand or safety rules?): PPE must be worn' in out['prompt'] and 'no unfinished walls' in out['prompt']
